"""The application tracker: one Excel workbook, one row per application.

    jhq tracker init
    jhq tracker add rows.json                 # list of {header: value}; IDs are assigned
    jhq tracker update updates.json           # list of {"id": "APP-001", "fields": {...}}
    jhq tracker event APP-001 interview [--date 2026-10-20] [--note "..."]
    jhq tracker search "acme"                 # company/title, case-insensitive
    jhq tracker show APP-001
    jhq tracker list [--json] [--cols "Company,Status"]
    jhq tracker age --days 30 [--apply]       # pending applications with no reply
    jhq tracker stats

Rules that keep the history honest:
- Every write backs the workbook up first (tracker/backups/, last 60 kept).
- The estimates and files recorded at application time are "protected": `update` refuses to
  change them unless --allow-protected is passed (only when the user asks for a correction).
- "Notes+" in update fields appends to Notes instead of replacing them.
- Dates: "YYYY-MM-DD", "today" or "yesterday" (in the configured timezone).
- If the workbook is open in Excel, writes stop with a clear message instead of failing later.
The workbook has no formulas or charts on purpose: the dashboard does the maths, and plain
cells survive any spreadsheet app.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import shutil
import time
from contextlib import contextmanager
from pathlib import Path

from . import config

STATUS = ["Applied", "On Hold", "Recruiter Screen", "Interview", "Final Interview", "Offer",
          "Rejected", "No Response", "Withdrawn"]
STAGE = ["Application Submitted", "Recruiter Screen", "First Interview", "Technical / Case",
         "Panel Interview", "Final Interview", "Offer", "Closed"]
OUTCOME = ["Pending", "Callback", "Interview", "Offer", "Rejected", "No Response", "Withdrawn"]
LISTS = {
    "STATUS": STATUS, "STAGE": STAGE, "OUTCOME": OUTCOME,
    "PRIORITY": ["High", "Normal", "Low"],
    "ARRANGEMENT": ["Remote", "Hybrid", "On-site", "Flexible", "-"],
    "CONTRACT": ["Permanent", "Fixed-term", "Contract", "Temporary", "Internship", "-"],
    "YESNO": ["Yes", "No"],
    "TAILOR": ["High", "Medium", "Light"],
    "DOMAIN": ["Yes", "Adjacent", "Project-based", "No"],
}
COLUMNS = [
    ("Application ID", 11, "text"), ("Company", 24, "text"), ("Job Title", 32, "text"),
    ("Status", 15, "list:STATUS"), ("Date Applied", 12, "date"), ("Status Updated", 12, "date"),
    ("Outcome Stage", 19, "list:STAGE"), ("Final Outcome", 13, "list:OUTCOME"),
    ("Priority", 9, "list:PRIORITY"), ("Location", 18, "text"),
    ("Work Arrangement", 14, "list:ARRANGEMENT"), ("Salary Range", 18, "text"),
    ("Salary Min", 11, "number"), ("Salary Max", 11, "number"), ("Currency", 8, "text"),
    ("Contract Type", 13, "list:CONTRACT"), ("Job Family", 20, "text"), ("Industry", 16, "text"),
    ("ATS Match When Applied", 13, "text"), ("Callback Estimate", 12, "text"),
    ("ATS Low", 8, "pct"), ("ATS High", 8, "pct"), ("Callback Low", 8, "pct"),
    ("Callback High", 8, "pct"), ("Keyword Coverage", 10, "pct"),
    ("Resume File", 40, "text"), ("Resume Pages", 8, "number"), ("Cover Letter", 8, "list:YESNO"),
    ("Tailoring Level", 10, "list:TAILOR"), ("Emphasis", 30, "text"),
    ("Direct Domain Experience", 12, "list:DOMAIN"), ("Main Strengths", 40, "text"),
    ("Main Gaps", 40, "text"), ("Resume Strategy", 40, "text"), ("Job Source / URL", 30, "text"),
    ("Follow-Up Date", 12, "date"), ("Callback Date", 12, "date"),
    ("First Interview Date", 12, "date"), ("Offer Date", 12, "date"), ("Rejection Date", 12, "date"),
    ("Notes", 50, "text"),
]
HEADERS = [c[0] for c in COLUMNS]
KIND = {c[0]: c[2] for c in COLUMNS}
PROTECTED = {"Application ID", "ATS Match When Applied", "Callback Estimate", "ATS Low", "ATS High",
             "Callback Low", "Callback High", "Keyword Coverage", "Resume File", "Resume Strategy"}
SHEET = "Applications"
KEEP_BACKUPS = 60

# What each event sets (dates use --date, else today).
EVENTS = {
    "screen": {"Status": "Recruiter Screen", "Outcome Stage": "Recruiter Screen",
               "Final Outcome": "Callback", "_date": "Callback Date"},
    "interview": {"Status": "Interview", "Outcome Stage": "First Interview",
                  "Final Outcome": "Interview", "_date": "First Interview Date", "_callback": True},
    "technical": {"Status": "Interview", "Outcome Stage": "Technical / Case",
                  "Final Outcome": "Interview", "_callback": True},
    "panel": {"Status": "Interview", "Outcome Stage": "Panel Interview",
              "Final Outcome": "Interview", "_callback": True},
    "final": {"Status": "Final Interview", "Outcome Stage": "Final Interview",
              "Final Outcome": "Interview", "_callback": True},
    "offer": {"Status": "Offer", "Outcome Stage": "Offer", "Final Outcome": "Offer",
              "_date": "Offer Date", "_callback": True},
    "rejected": {"Status": "Rejected", "Outcome Stage": "Closed", "Final Outcome": "Rejected",
                 "_date": "Rejection Date"},
    "no-response": {"Status": "No Response", "Outcome Stage": "Closed", "Final Outcome": "No Response"},
    "withdrawn": {"Status": "Withdrawn", "Outcome Stage": "Closed", "Final Outcome": "Withdrawn"},
    "on-hold": {"Status": "On Hold"},
}


class TrackerError(SystemExit):
    pass


# ---------------------------------------------------------------- file handling
def _path(paths=None) -> Path:
    return (paths or config.Paths()).tracker


def excel_has_it_open(path: Path) -> bool:
    """Excel (Windows and Mac) and LibreOffice leave a lock file next to an open workbook."""
    lock_names = [f"~${path.name}", f".~lock.{path.name}#"]
    if any((path.parent / n).exists() for n in lock_names):
        return True
    if os.name == "nt" and path.exists():
        try:
            with open(path, "r+b"):
                pass
        except PermissionError:
            return True
    return False


@contextmanager
def write_lock(path: Path, timeout=30):
    """One writer at a time (Claude and the dashboard sync can both write)."""
    lock = path.with_name(path.name + ".writing")
    start = time.time()
    while True:
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
            break
        except FileExistsError:
            if time.time() - lock.stat().st_mtime > 300:  # stale lock from a crashed run
                lock.unlink(missing_ok=True)
                continue
            if time.time() - start > timeout:
                raise TrackerError("Another tracker write is in progress; try again in a minute.")
            time.sleep(0.5)
    try:
        yield
    finally:
        lock.unlink(missing_ok=True)


def backup(path: Path, paths=None) -> str | None:
    if not path.exists():
        return None
    bdir = (paths or config.Paths()).backups
    bdir.mkdir(parents=True, exist_ok=True)
    out = bdir / f"applications_{dt.datetime.now():%Y-%m-%d_%H%M%S_%f}.xlsx"
    shutil.copy2(path, out)
    old = sorted(bdir.glob("applications_*.xlsx"))
    for f in old[:-KEEP_BACKUPS]:
        f.unlink(missing_ok=True)
    return str(out)


def init(paths=None, force=False) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = _path(paths)
    if path.exists() and not force:
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    head_font, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F4FC4")
    for i, (name, width, kind) in enumerate(COLUMNS, 1):
        c = ws.cell(1, i)
        c.font, c.fill = head_font, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(HEADERS)).column_letter}1"

    lists = wb.create_sheet("Lists")
    for ci, (key, values) in enumerate(LISTS.items(), 1):
        lists.cell(1, ci, key).font = Font(bold=True)
        for ri, v in enumerate(values, 2):
            lists.cell(ri, ci, v)
        col = lists.cell(1, ci).column_letter
        for i, (name, _, kind) in enumerate(COLUMNS, 1):
            if kind == f"list:{key}":
                letter = ws.cell(1, i).column_letter
                dv = DataValidation(type="list", formula1=f"=Lists!${col}$2:${col}${len(values) + 1}",
                                    allow_blank=True, showErrorMessage=False)
                dv.add(f"{letter}2:{letter}2000")
                ws.add_data_validation(dv)
    lists.sheet_state = "hidden"

    guide = wb.create_sheet("Read Me")
    guide["A1"] = "Job Hunt HQ application tracker"
    guide["A1"].font = Font(bold=True, size=14)
    for r, line in enumerate([
        "One row per application. Claude updates it for you; you can also edit it by hand.",
        "Close this file in Excel before asking Claude to update it (Excel locks open files).",
        "Estimates (ATS, callback) are recorded when you apply and are never revised later,",
        "so you can compare what was expected with what happened.",
        "Backups of every change: the 'backups' folder next to this file.",
    ], 3):
        guide.cell(r, 1, line)
    guide.column_dimensions["A"].width = 100
    wb.save(path)
    return path


def _load(path: Path, read_only=False):
    import openpyxl
    if not path.exists():
        raise TrackerError(f"No tracker at {path}. Run `jhq tracker init`.")
    blob = path.read_bytes()  # read, then let go of the file at once
    return openpyxl.load_workbook(io.BytesIO(blob), read_only=read_only, data_only=True)


def _fmt(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def rows(paths=None) -> list[dict]:
    path = _path(paths)
    wb = _load(path, read_only=True)
    try:
        grid = list(wb[SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not grid:
        return []
    head = [h for h in grid[0]]
    out = []
    for r in grid[1:]:
        if not r or not r[0]:
            continue
        out.append({h: _fmt(v) for h, v in zip(head, r) if h})
    return out


# ---------------------------------------------------------------- value handling
def _coerce(header, value):
    kind = KIND.get(header)
    if kind is None:
        raise TrackerError(f"Unknown column '{header}'. Columns: {', '.join(HEADERS)}")
    if value is None or value == "":
        return None
    if kind == "date":
        try:
            return config.parse_day(value)
        except ValueError as e:
            raise TrackerError(f"{header}: '{value}' is not a date (use YYYY-MM-DD, today, yesterday)") from e
    if kind == "pct":
        if isinstance(value, str):
            v = value.strip().rstrip("%")
            value = float(v) / (100 if value.strip().endswith("%") or float(v) > 1 else 1)
        value = float(value)
        if value > 1:
            value /= 100
        return round(value, 4)
    if kind == "number":
        if isinstance(value, str):
            s = value.replace(",", "").replace("$", "").strip()
            try:
                return float(s) if "." in s else int(s)
            except ValueError as e:
                raise TrackerError(f"{header}: '{value}' is not a number") from e
        return value
    if kind.startswith("list:"):
        allowed = LISTS[kind[5:]]
        match = next((a for a in allowed if a.lower() == str(value).strip().lower()), None)
        if not match:
            raise TrackerError(f"{header}: '{value}' is not one of {allowed}")
        return match
    return str(value)


def _write_cell(ws, r, header, value):
    c = ws.cell(r, HEADERS.index(header) + 1, value)
    kind = KIND[header]
    if kind == "date" and value is not None:
        c.number_format = "yyyy-mm-dd"
    elif kind == "pct" and value is not None:
        c.number_format = "0%"


def _header_index(ws):
    head = [c.value for c in ws[1]]
    missing = [h for h in HEADERS if h not in head]
    if missing:
        raise TrackerError(f"The tracker is missing columns {missing}. Was it edited by hand? "
                           "Restore a backup or add them back.")
    return head


def _save(wb, path):
    tmp = path.with_name(path.stem + ".saving.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


def _guard(path):
    if excel_has_it_open(path):
        raise TrackerError(f"The tracker is open in Excel ({path.name}). Close it, then try again.")


# ---------------------------------------------------------------- writes
def add(new_rows: list[dict], paths=None) -> dict:
    path = _path(paths)
    init(paths)
    _guard(path)
    with write_lock(path):
        wb = _load(path)
        ws = wb[SHEET]
        _header_index(ws)
        existing = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
        nums = [int(str(x).split("-")[-1]) for x in existing if str(x).startswith("APP-")
                and str(x).split("-")[-1].isdigit()]
        nxt = max(nums, default=0) + 1
        written = []
        prepared = []
        for raw in new_rows:
            rec = {k: v for k, v in raw.items() if k != "Application ID"}
            if not rec.get("Company") or not rec.get("Job Title"):
                raise TrackerError("Each new row needs Company and Job Title.")
            rec.setdefault("Status", "Applied")
            rec.setdefault("Outcome Stage", "Application Submitted")
            rec.setdefault("Final Outcome", "Pending")
            rec.setdefault("Priority", "Normal")
            if rec.get("Date Applied") and not rec.get("Status Updated"):
                rec["Status Updated"] = rec["Date Applied"]
            prepared.append({h: _coerce(h, v) for h, v in rec.items()})  # validate all first
        bk = backup(path, paths)
        r = ws.max_row + 1
        while r > 2 and not ws.cell(r - 1, 1).value:
            r -= 1
        for rec in prepared:
            app_id = f"APP-{nxt:03d}"
            nxt += 1
            _write_cell(ws, r, "Application ID", app_id)
            for h, v in rec.items():
                _write_cell(ws, r, h, v)
            written.append({"Application ID": app_id, **{h: _fmt(v) for h, v in rec.items()}})
            r += 1
        _save(wb, path)
    return {"added": written, "backup": bk, "tracker": str(path)}


def update(updates: list[dict], paths=None, allow_protected=False) -> dict:
    path = _path(paths)
    _guard(path)
    with write_lock(path):
        wb = _load(path)
        ws = wb[SHEET]
        _header_index(ws)
        index = {str(ws.cell(r, 1).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
        plan = []
        for u in updates:
            app_id, fields = str(u.get("id", "")), dict(u.get("fields") or {})
            if app_id not in index:
                raise TrackerError(f"No row {app_id}. Search first: jhq tracker search <company>")
            blocked = sorted(set(fields) & PROTECTED)
            if blocked and not allow_protected:
                raise TrackerError(f"{app_id}: {blocked} were recorded at application time and are "
                                   "protected. Use --allow-protected only if the user asked to correct them.")
            note_add = fields.pop("Notes+", None)
            coerced = {h: _coerce(h, v) for h, v in fields.items()}
            plan.append((app_id, index[app_id], coerced, note_add))
        bk = backup(path, paths)
        out = []
        for app_id, r, coerced, note_add in plan:
            for h, v in coerced.items():
                _write_cell(ws, r, h, v)
            if note_add:
                ci = HEADERS.index("Notes") + 1
                cur = ws.cell(r, ci).value
                ws.cell(r, ci, f"{cur}; {note_add}" if cur else note_add)
            out.append({"id": app_id, **{h: _fmt(v) for h, v in coerced.items()},
                        **({"Notes+": note_add} if note_add else {})})
        _save(wb, path)
    return {"updated": out, "backup": bk, "tracker": str(path)}


def event(app_id: str, name: str, date=None, note=None, paths=None) -> dict:
    """Apply a status event with the standard field changes. Never clears earlier dates."""
    if name not in EVENTS:
        raise TrackerError(f"Unknown event '{name}'. Use one of: {', '.join(EVENTS)}")
    spec = EVENTS[name]
    day = config.parse_day(date) if date else config.today()
    current = next((r for r in rows(paths) if r.get("Application ID") == app_id), None)
    if not current:
        raise TrackerError(f"No row {app_id}.")
    fields = {k: v for k, v in spec.items() if not k.startswith("_")}
    fields["Status Updated"] = day.isoformat()
    if spec.get("_date") and (date or not current.get(spec["_date"])):
        fields[spec["_date"]] = day.isoformat()
    if spec.get("_callback") and not current.get("Callback Date"):
        fields["Callback Date"] = day.isoformat()
    if name == "on-hold":
        fields.pop("Outcome Stage", None)
    if note:
        fields["Notes+"] = note
    return update([{"id": app_id, "fields": fields}], paths)


def aging(days: int, paths=None) -> tuple[list[dict], list[dict]]:
    """(eligible, skipped_without_date) for a no-response batch."""
    t = config.today()
    eligible, skipped = [], []
    for r in rows(paths):
        if r.get("Status") not in ("Applied", "On Hold") or r.get("Final Outcome") not in (None, "Pending"):
            continue
        if any(r.get(k) for k in ("Callback Date", "First Interview Date", "Offer Date", "Rejection Date")):
            continue
        if not r.get("Date Applied"):
            skipped.append(r)
            continue
        age = (t - dt.date.fromisoformat(r["Date Applied"])).days
        if age > days:
            eligible.append({**r, "_age": age})
    return eligible, skipped


def stats(paths=None) -> dict:
    rs = rows(paths)
    by = lambda k: {v: sum(1 for r in rs if (r.get(k) or "-") == v)  # noqa: E731
                    for v in sorted({r.get(k) or "-" for r in rs})}
    t = config.today()
    week = sum(1 for r in rs if r.get("Date Applied")
               and (t - dt.date.fromisoformat(r["Date Applied"])).days < 7)
    return {"total": len(rs), "last_7_days": week, "by_status": by("Status"),
            "by_final_outcome": by("Final Outcome"), "missing_date_applied":
            [r["Application ID"] for r in rs if not r.get("Date Applied")]}


# ---------------------------------------------------------------- CLI
def _load_json(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


def main(argv=None):
    config.utf8_console()
    ap = argparse.ArgumentParser(prog="jhq tracker", description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init")
    s = sub.add_parser("add"); s.add_argument("json")
    s = sub.add_parser("update"); s.add_argument("json"); s.add_argument("--allow-protected", action="store_true")
    s = sub.add_parser("event"); s.add_argument("id"); s.add_argument("name", choices=list(EVENTS))
    s.add_argument("--date"); s.add_argument("--note")
    s = sub.add_parser("search"); s.add_argument("query")
    s = sub.add_parser("show"); s.add_argument("id")
    s = sub.add_parser("list"); s.add_argument("--json", action="store_true"); s.add_argument("--cols")
    s = sub.add_parser("age"); s.add_argument("--days", type=int, required=True)
    s.add_argument("--apply", action="store_true")
    sub.add_parser("stats")
    a = ap.parse_args(argv)

    if a.cmd == "init":
        print(init())
        return
    if a.cmd == "add":
        data = _load_json(a.json)
        print(json.dumps(add(data if isinstance(data, list) else [data]), indent=2, ensure_ascii=False))
        return
    if a.cmd == "update":
        data = _load_json(a.json)
        print(json.dumps(update(data if isinstance(data, list) else [data],
                                allow_protected=a.allow_protected), indent=2, ensure_ascii=False))
        return
    if a.cmd == "event":
        print(json.dumps(event(a.id, a.name, a.date, a.note), indent=2, ensure_ascii=False))
        return
    if a.cmd in ("search", "show", "list"):
        rs = rows()
        if a.cmd == "search":
            q = a.query.lower()
            rs = [r for r in rs if q in str(r.get("Company", "")).lower() or q in str(r.get("Job Title", "")).lower()]
        if a.cmd == "show":
            rs = [r for r in rs if r.get("Application ID") == a.id]
        if getattr(a, "cols", None):
            keep = [c.strip() for c in a.cols.split(",")]
            rs = [{k: r.get(k) for k in ["Application ID", *keep]} for r in rs]
        if a.cmd == "list" and not a.json:
            for r in rs:
                print(" | ".join(str(r.get(k, "")) for k in ("Application ID", "Company", "Job Title",
                                                              "Status", "Date Applied")))
        else:
            print(json.dumps([{k: v for k, v in r.items() if v not in (None, "")} for r in rs],
                             indent=2, ensure_ascii=False))
        print(f"# {len(rs)} row(s)")
        return
    if a.cmd == "age":
        eligible, skipped = aging(a.days)
        print(f"{len(eligible)} pending application(s) older than {a.days} days with no reply:")
        for r in eligible:
            print(f"  {r['Application ID']} | {r['Company']} | {r['Job Title']} | applied {r['Date Applied']} ({r['_age']}d)")
        if skipped:
            print(f"{len(skipped)} skipped (no Date Applied): " + ", ".join(r["Application ID"] for r in skipped))
        if a.apply and eligible:
            note = f"Aged out after {a.days}+ days with no response"
            res = update([{"id": r["Application ID"], "fields": {
                "Status": "No Response", "Outcome Stage": "Closed", "Final Outcome": "No Response",
                "Status Updated": "today", "Notes+": note}} for r in eligible])
            print(f"Updated {len(res['updated'])} row(s). Backup: {res['backup']}")
        elif eligible:
            print("Nothing changed. Show this list to the user; rerun with --apply after they say yes.")
        return
    if a.cmd == "stats":
        print(json.dumps(stats(), indent=2))


if __name__ == "__main__":
    main()
