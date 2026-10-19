"""Turn a simple Markdown file into a clean .docx (interview prep packs, learning plans).

    jhq md2docx input.md output.docx [--overwrite]

Handles # / ## / ### headings, - or * bullets (2-space indent = level 2), 1. numbered lists,
| tables |, **bold**, [label](https://url) links, > quotes, --- rules and blank lines.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from . import build_package as bp
from . import config

SIZE = 10.5


def convert(md_text, out):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = bp.FONT, Pt(SIZE)
    md_text = re.sub(r"`([^`\n]+)`", r"\1", md_text)
    md_text = re.sub(r"(?<![*\w])\*([^*\s][^*\n]*?)\*(?![*\w])", r"\1", md_text)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            bp.set_para(p, before=(0 if level == 1 else 10), after=4)
            bp.set_font(p.add_run(m.group(2).replace("**", "")), {1: 16, 2: 12.5, 3: 11}[level], True)
            if level == 2:
                bp.add_rule(p)
            p.paragraph_format.keep_with_next = True
        elif re.match(r"^-{3,}$", ln.strip()):
            bp.add_rule(doc.add_paragraph())
        elif ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.paragraphs[0].text = ""
                        text = row[ci] if ci < len(row) else ""
                        if ri == 0 and text:
                            text = f"**{text.replace('**', '')}**"
                        bp.runs_from_markup(cell.paragraphs[0], text, SIZE - 1)
                doc.add_paragraph()
            continue
        elif re.match(r"^\s*([-*]|\d+\.)\s+", ln):
            indent = len(ln) - len(ln.lstrip())
            numbered = bool(re.match(r"^\s*\d+\.", ln))
            text = re.sub(r"^\s*([-*]|\d+\.)\s+", "", ln)
            p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            bp.set_para(p, after=2, line=1.05)
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (indent // 2))
            bp.runs_from_markup(p, text, SIZE)
        elif ln.startswith(">"):
            p = doc.add_paragraph()
            bp.set_para(p, after=4, line=1.05)
            p.paragraph_format.left_indent = Inches(0.3)
            bp.runs_from_markup(p, ln.lstrip("> "), SIZE)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        else:
            p = doc.add_paragraph()
            bp.set_para(p, after=5, line=1.08)
            bp.runs_from_markup(p, ln, SIZE)
        i += 1
    doc.save(out)


def main(argv=None):
    config.utf8_console()
    ap = argparse.ArgumentParser(prog="jhq md2docx")
    ap.add_argument("md")
    ap.add_argument("docx")
    ap.add_argument("--overwrite", action="store_true")
    a = ap.parse_args(argv)
    out = Path(a.docx)
    if out.exists() and not a.overwrite:
        raise SystemExit(f"Refusing to overwrite {out} (pass --overwrite for a file made this session)")
    out.parent.mkdir(parents=True, exist_ok=True)
    convert(Path(a.md).read_text(encoding="utf-8"), out)
    print(out)


if __name__ == "__main__":
    main()
