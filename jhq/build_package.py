"""Build a job-application package, then check it.

    jhq build spec.json [--overwrite] [--qa-dir DIR] [--qa-only]

Writes into <data>/applications/<Company> - <Job Title>/:
    Job Description.docx              summary + the original posting, verbatim
    <Name>_Resume.docx / .pdf         tailored resume
    <Name>_Cover_Letter.docx / .pdf   cover letter (if the spec has one)
    tracker_meta.json                 final estimates, read when the application is tracked

Name, contact line, links, font and "never write" phrases come from profile/profile.json,
so the spec only holds content. Then QA runs: layout, links, bold count, page fill, a
measured ATS check (see ats_check.py) and cover-letter checks. It prints a JSON report.

spec.json (any of resume / jd / cover_letter may be left out):
{
  "company": "Acme Analytics",            # with "title", decides the folder name
  "title": "Data Analyst",
  "link": "https://...",                  # optional: posting URL, saved to tracker_meta.json
  "resume": {
    "location": null,                     # null = profile location
    "headline": null,                     # optional line under the name
    "sizes": {...},                       # optional overrides of DEFAULT_SIZES
    "summary": "Data analyst with **3+ years** ...",          # **bold** markup
    "experience": [
      {"title": "Data Analyst | Example Co - Chicago, IL | Jan 2022 - Present",
       "groups": [{"subhead": "Reporting:", "bullets": ["Built **dashboards** ..."]}]}
    ],
    "skills": [["Analytics", "SQL, Python, Tableau"]],
    "education": [{"degree": "B.S. Statistics", "detail": "Example University | 2016 - 2020"}],
    "coursework": null, "certifications": [], "languages": null,
    "projects": {"heading": "Projects", "bullets": ["[Name](https://url) - what it does"]},
    "projects_after_experience": false,
    "extra_sections": [{"heading": "Volunteering", "bullets": ["..."]}],
    "order": null                         # e.g. ["summary","experience","projects","skills","education"]
  },
  "jd": {"summary": [["Company", "..."], ["Job title", "..."], ["Location", "..."],
                     ["Work arrangement", "..."], ["Salary", "Not stated"],
                     ["Most relevant skills", "..."], ["Nice-to-have skills", "..."],
                     ["Experience asked", "..."]],
         "text_file": "jd.txt"},          # the posting text, copied verbatim
  "ats": {"required": ["SQL", {"term": "stakeholder management", "aka": ["stakeholders"]}],
          "preferred": ["dbt"]},
  "cover_letter": {"date": null, "recipient": ["Hiring Team", "Acme Analytics", "Chicago, IL"],
                   "salutation": "Dear Hiring Team,", "paragraphs": ["..."],
                   "bullets": null, "bullets_after": 1, "closing": "Sincerely,", "size": 11},
  "tracker": {"fields": {"ATS Match When Applied": "78-85%", "Callback Estimate": "12-18%", ...}}
}
A spec with only company/title + cover_letter adds a letter to an existing package.
"""
from __future__ import annotations

import argparse
import json
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import config
from .pdf_export import ExportError, export_pdf

FONT = "Calibri"
DEFAULT_SIZES = {"name": 18, "headline": 10, "contact": 9.5, "heading": 11, "body": 10,
                 "role": 10.2, "subhead": 10, "bullet": 9.8, "skills": 9.6, "line": 1.0,
                 "bullet_line": 1.0, "bullet_after": 1.5, "heading_before": 6,
                 "role_before": 4, "subhead_before": 2,
                 "margins": [0.5, 0.5, 0.6, 0.6]}
DEFAULT_ORDER = ["summary", "experience", "skills", "education", "certifications",
                 "languages", "projects", "extra"]


# ---------- docx helpers ----------
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean(text):
    """Drop invisible control characters: Word files can't contain them (job postings sometimes do)."""
    return CONTROL.sub("", str(text))


def set_font(run, size=None, bold=None, font=None):
    font = font or FONT
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rpr.rFonts.set(qn(a), font)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_para(p, before=0, after=0, line=1.0, align=None):
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    if align is not None:
        p.alignment = align


MARKUP = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)\s]+\))")
LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")


def runs_from_markup(p, text, size, allow_bold=True):
    """**bold** and [label](https://url) markup; links become real, clickable hyperlinks."""
    for part in MARKUP.split(text):
        if not part:
            continue
        m = LINK.fullmatch(part)
        if m:
            add_hyperlink(p, m.group(1), m.group(2), size)
            continue
        bold = part.startswith("**") and part.endswith("**")
        set_font(p.add_run(clean(part[2:-2] if bold else part)), size, bold and allow_bold)


def add_hyperlink(p, text, url, size):
    r_id = p.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), FONT)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(round(size * 2))))
    for el in (fonts, color, u, sz):
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = clean(text)
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def add_rule(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def setup(doc, size, margins):
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = [Inches(m) for m in margins]
    for name in ("Normal", "List Bullet"):
        st = doc.styles[name]
        st.font.name = FONT
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)


def heading(doc, text, sz):
    p = doc.add_paragraph()
    set_para(p, before=sz["heading_before"], after=2)
    set_font(p.add_run(clean(text.upper())), sz["heading"], True)
    add_rule(p)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, size, sz, allow_bold=True):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, after=sz["bullet_after"], line=sz["bullet_line"])
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    runs_from_markup(p, text, size, allow_bold)
    return p


def plain(doc, chunks, size, before=0, after=0, line=1.0, align=None):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=line, align=align)
    for text, bold in chunks:
        set_font(p.add_run(clean(text)), size, bold)
    return p


def contact_line(doc, prof, location, size, after=2):
    """'City | phone | email | LinkedIn | Portfolio' with real links. Phone and email stay plain
    text so an ATS can read them."""
    bits = [b for b in (location or prof.get("location"), prof.get("phone"), prof.get("email")) if b]
    p = plain(doc, [(" | ".join(bits), False)], size, after=after, align=WD_ALIGN_PARAGRAPH.CENTER)
    for link in prof.get("links") or []:
        set_font(p.add_run(" | "), size)
        add_hyperlink(p, link["label"], link["url"], size)
    return p


# ---------- builders ----------
def build_resume(spec, prof, path):
    sz = {**DEFAULT_SIZES, **(spec.get("sizes") or {})}
    doc = Document()
    setup(doc, sz["body"], sz["margins"])
    C = WD_ALIGN_PARAGRAPH.CENTER

    plain(doc, [(prof["name"], True)], sz["name"], align=C)
    if spec.get("headline"):
        plain(doc, [(spec["headline"], True)], sz["headline"], align=C)
    contact_line(doc, prof, spec.get("location"), sz["contact"])

    def summary():
        if spec.get("summary"):
            heading(doc, spec.get("summary_heading", "Summary"), sz)
            p = doc.add_paragraph()
            set_para(p, after=1, line=sz["line"])
            runs_from_markup(p, spec["summary"], sz["body"])

    def experience():
        if not spec.get("experience"):
            return
        heading(doc, spec.get("experience_heading", "Professional Experience"), sz)
        for role in spec["experience"]:
            plain(doc, [(role["title"], True)], sz["role"], before=sz["role_before"], after=1)
            for g in role.get("groups", []):
                if g.get("subhead"):
                    plain(doc, [(g["subhead"], True)], sz["subhead"], before=sz["subhead_before"], after=1)
                for b in g["bullets"]:
                    bullet(doc, b, sz["bullet"], sz)

    def skills():
        if not spec.get("skills"):
            return
        heading(doc, spec.get("skills_heading", "Skills"), sz)
        for label, items in spec["skills"]:
            p = doc.add_paragraph(style="List Bullet")
            set_para(p, after=sz["bullet_after"], line=sz["bullet_line"])
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            set_font(p.add_run(clean(f"{label}: ")), sz["skills"], True)
            set_font(p.add_run(clean(items.replace("**", ""))), sz["skills"], False)

    def education():
        if not spec.get("education"):
            return
        heading(doc, "Education", sz)
        for i, e in enumerate(spec["education"]):
            plain(doc, [(e["degree"], True), (f" - {e['detail']}" if e.get("detail") else "", False)],
                  sz["body"], after=1)
            if i == 0 and spec.get("coursework"):
                plain(doc, [("Relevant coursework: ", True), (spec["coursework"], False)], sz["skills"])

    def certifications():
        if spec.get("certifications"):
            heading(doc, "Certifications", sz)
            plain(doc, [(" | ".join(spec["certifications"]), False)], sz["skills"])

    def languages():
        if spec.get("languages"):
            heading(doc, "Languages", sz)
            plain(doc, [(spec["languages"], False)], sz["skills"])

    def projects():
        pr = spec.get("projects")
        if pr and pr.get("bullets"):
            heading(doc, pr.get("heading", "Projects"), sz)
            for b in pr["bullets"]:
                bullet(doc, b, sz["bullet"], sz, allow_bold=False)

    def extra():
        for sec in spec.get("extra_sections") or []:
            heading(doc, sec["heading"], sz)
            for b in sec.get("bullets", []):
                bullet(doc, b, sz["bullet"], sz, allow_bold=False)

    parts = {"summary": summary, "experience": experience, "skills": skills, "education": education,
             "certifications": certifications, "languages": languages, "projects": projects,
             "extra": extra}
    order = list(spec.get("order") or DEFAULT_ORDER)
    if not spec.get("order") and spec.get("projects_after_experience"):
        order.remove("projects")
        order.insert(order.index("experience") + 1, "projects")
    for key in order:
        if key not in parts:
            raise SystemExit(f"Unknown section in 'order': {key}. Use: {', '.join(parts)}")
        parts[key]()
    doc.save(path)


def build_jd(spec, title, path, spec_dir):
    sz = {**DEFAULT_SIZES, "heading": 12, "heading_before": 8}
    doc = Document()
    setup(doc, 10.5, [0.7, 0.7, 0.8, 0.8])
    plain(doc, [(title, True)], 16, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc, "Job Summary", sz)
    for label, value in spec.get("summary") or []:
        plain(doc, [(f"{label}: ", True), (value or "-", False)], 10.5, after=2)
    heading(doc, "Original Job Description", sz)
    tf = Path(spec["text_file"])
    if not tf.is_absolute():
        tf = spec_dir / tf
    raw = tf.read_text(encoding="utf-8")
    for line in raw.splitlines():  # verbatim; blank lines kept
        p = doc.add_paragraph()
        set_para(p, after=2)
        if line:
            set_font(p.add_run(clean(line)), 10.5)
    doc.save(path)


def build_cover(spec, prof, path):
    size = spec.get("size", 11)
    doc = Document()
    setup(doc, size, spec.get("margins", [0.75, 0.75, 0.9, 0.9]))
    plain(doc, [(prof["name"], True)], 18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rule(contact_line(doc, prof, spec.get("location"), 9.5))
    date = spec.get("date") or config.today().strftime("%B %d, %Y").replace(" 0", " ")
    plain(doc, [(date, False)], size, before=14, after=10)
    for line in spec.get("recipient") or []:
        plain(doc, [(line, False)], size)
    plain(doc, [(spec.get("salutation", "Dear Hiring Team,"), False)], size, before=10, after=8)
    for i, para in enumerate(spec["paragraphs"]):
        p = doc.add_paragraph()
        set_para(p, after=8, line=1.05)
        runs_from_markup(p, para, size)
        if spec.get("bullets") and i == spec.get("bullets_after", 1):
            for b in spec["bullets"]:
                bp = bullet(doc, b, size, {**DEFAULT_SIZES, "bullet_after": 3, "bullet_line": 1.05})
                bp.paragraph_format.left_indent = Inches(0.35)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
    plain(doc, [(spec.get("closing", "Sincerely,"), False)], size, before=4, after=2)
    plain(doc, [(prof["name"], False)], size)
    doc.save(path)


# ---------- QA ----------
CLICHES = ["passionate about", "great fit", "perfect fit", "leverage my", "hit the ground running",
           "drive results", "results-driven", "synerg", "to whom it may concern", "team player",
           "think outside the box", "dynamic environment", "i believe i would", "detail-oriented",
           "fast-paced environment", "go-getter", "excited to apply", "i am writing to express"]


def _profile_links(prof):
    return [l["url"] for l in prof.get("links") or []]


def _forbidden(prof):
    return [f for f in prof.get("forbidden_phrases") or [] if f]


def short_last_lines(pdf, max_words=2):
    """Paragraphs whose last line holds only 1-2 words (a wasted line that looks sloppy)."""
    found = []
    for page in pdf:
        for block in page.get_text("dict")["blocks"]:
            lines = [(ln["bbox"][0], " ".join(s["text"] for s in ln["spans"]).strip())
                     for ln in block.get("lines", [])]
            lines = [(x, t) for x, t in lines if t]
            for i in range(1, len(lines)):
                x, text = lines[i]
                # A paragraph ends where the block ends or the next line starts further left
                # (the next bullet). Continuation lines of a bullet are indented past its bullet.
                ends = i == len(lines) - 1 or lines[i + 1][0] < x - 3
                if ends and len(text.split()) <= max_words and not text.isupper():
                    found.append(text)
    return found


def qa_resume(docx_path, pdf_path, qa_dir, prof, extra_links):
    import pymupdf
    style = prof.get("style") or {}
    report, warn = {}, []
    doc = Document(docx_path)

    rels = [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    report["docx_links"] = rels
    for u in (*_profile_links(prof), *extra_links):
        if u not in rels:
            warn.append(f"missing DOCX hyperlink {u}")

    paras = doc.paragraphs
    headings = [p for p in paras if p._p.pPr is not None and p._p.pPr.find(qn("w:pBdr")) is not None]
    report["headings_with_rule"] = [p.text for p in headings]
    for p in headings:
        if p.text != p.text.upper() or not all(r.bold for r in p.runs):
            warn.append(f"heading not uppercase/bold: {p.text}")

    bold_frags, section = [], None
    for p in paras:
        if p in headings:
            section = p.text
            continue
        if section in ("SUMMARY", "PROFESSIONAL EXPERIENCE", "EXPERIENCE") and p.runs:
            if all(r.bold for r in p.runs if r.text.strip()):
                continue  # role titles and subheads are fully bold by design
            bold_frags += [r.text for r in p.runs if r.bold and r.text.strip()]
    report["bold_fragments"] = bold_frags
    max_bold = style.get("max_bold", 10)
    if len(bold_frags) > max_bold:
        warn.append(f"{len(bold_frags)} bold fragments in summary/experience (target <= {max_bold}); "
                    "bold only the results a recruiter should see first")

    in_skills = False
    for p in paras:
        if p in headings:
            in_skills = p.text in ("SKILLS", "TECHNICAL SKILLS")
            continue
        if in_skills and p.runs:
            if [r.text for r in p.runs[1:] if r.bold] or not p.runs[0].bold:
                warn.append(f"skills line bolding wrong (bold the label only): {p.text[:50]}")

    pdf = pymupdf.open(pdf_path)
    text = "\n".join(pg.get_text() for pg in pdf)
    report["pages"] = len(pdf)
    for f in _forbidden(prof):
        if f.lower() in text.lower():
            warn.append(f"forbidden text present: {f}")
    pdf_links = [l.get("uri") for pg in pdf for l in pg.get_links() if l.get("uri")]
    report["pdf_links"] = pdf_links
    for u in (*_profile_links(prof), *extra_links):
        if not any(u.rstrip("/") in (x or "") for x in pdf_links):
            warn.append(f"PDF missing link {u}")

    orphans = short_last_lines(pdf)
    report["short_last_lines"] = orphans
    if orphans:
        warn.append(f"{len(orphans)} bullet(s) end with 1-2 words on their own line: {orphans[:4]}; "
                    "reword to fit one line or fill the second line")

    last = pdf[-1]
    usable = last.rect.height - doc.sections[0].bottom_margin.pt
    ys = [b[3] for b in last.get_text("blocks") if b[4].strip()]
    fill = round(100 * max(ys) / usable, 1) if ys else 0
    report["last_page_fill_pct"] = fill
    fill_min = style.get("page_fill_min", 92)
    if len(pdf) == 1 and fill < fill_min:
        warn.append(f"one page only {fill}% filled (target {fill_min}-99%): grow sizes/spacing "
                    "a little or add a relevant, true bullet")
    if len(pdf) > 1 and fill < 45:
        warn.append(f"last page only {fill}% filled: tighten to 1 page or fill page 2 with real content")
    if len(pdf) > 2:
        warn.append(f"{len(pdf)} pages: too long")

    qa_dir.mkdir(parents=True, exist_ok=True)
    pngs = []
    for i, pg in enumerate(pdf):
        out = qa_dir / f"page{i + 1}.png"
        pg.get_pixmap(dpi=110).save(str(out))
        pngs.append(str(out))
    pdf.close()
    report["renders"] = pngs
    report["warnings"] = warn
    return report


def qa_cover(spec, pdf_path, qa_dir, prof, company):
    import pymupdf
    warn = []
    body = "\n".join(spec["paragraphs"] + (spec.get("bullets") or []))
    body_plain = re.sub(r"\*\*|\[([^\]]+)\]\([^)]+\)", r"\1", body)
    words = len(body_plain.split())
    rep = {"words": words}
    if words > 420:
        warn.append(f"{words} words: too long (target 280-380, one page)")
    elif words < 220:
        warn.append(f"{words} words: thin (target 280-380)")
    if re.search("[—–]", body):
        warn.append("em/en dash in the letter: use commas or periods")
    low = body_plain.lower()
    hits = [c for c in CLICHES if c in low]
    if hits:
        warn.append(f"cliches: {hits}")
    if company and company.lower() not in low:
        warn.append(f"company name '{company}' never appears in the body")
    for f in _forbidden(prof):
        if f.lower() in low:
            warn.append(f"forbidden text present: {f}")
    pdf = pymupdf.open(pdf_path)
    rep["pages"] = len(pdf)
    if len(pdf) != 1:
        warn.append(f"cover letter is {len(pdf)} pages: must be exactly 1")
    links = [l.get("uri") for pg in pdf for l in pg.get_links() if l.get("uri")]
    for u in _profile_links(prof):
        if not any(u.rstrip("/") in (x or "") for x in links):
            warn.append(f"cover PDF missing link {u}")
    qa_dir.mkdir(parents=True, exist_ok=True)
    out = qa_dir / "cover_page1.png"
    pdf[0].get_pixmap(dpi=110).save(str(out))
    pdf.close()
    rep["renders"] = [str(out)]
    rep["warnings"] = warn
    return rep


# ---------- package ----------
def package_folder(spec, paths) -> Path:
    if spec.get("folder"):
        folder = Path(spec["folder"])
        if not folder.is_absolute():
            folder = paths.applications / folder
    else:
        if not (spec.get("company") and spec.get("title")):
            raise SystemExit("The spec needs 'company' and 'title' (or 'folder').")
        folder = paths.applications / config.safe_name(f"{spec['company']} - {spec['title']}")
    # Never write outside the applications folder, whatever the spec says.
    root = paths.applications.resolve()
    if root not in (folder.resolve(), *folder.resolve().parents):
        raise SystemExit(f"Refusing to write outside {root}: {folder}")
    return folder


def file_names(folder, prof):
    stub = config.file_stub(prof["name"])
    return {"resume_docx": folder / f"{stub}_Resume.docx", "resume_pdf": folder / f"{stub}_Resume.pdf",
            "cover_docx": folder / f"{stub}_Cover_Letter.docx",
            "cover_pdf": folder / f"{stub}_Cover_Letter.pdf",
            "jd_docx": folder / "Job Description.docx", "meta": folder / "tracker_meta.json"}


def write_meta(spec, folder, files, rep):
    meta_path = files["meta"]
    meta = config.read_json(meta_path, {}) or {}
    tr = spec.get("tracker") or {}
    meta.update({"company": spec.get("company") or meta.get("company"),
                 "title": spec.get("title") or meta.get("title"),
                 "built": config.today().isoformat()})
    if spec.get("link"):
        meta["link"] = spec["link"]
    fields = {**(meta.get("fields") or {}), **(tr.get("fields") or {})}
    if files["resume_pdf"].exists():
        fields["Resume File"] = str(files["resume_pdf"])
        if rep.get("pages"):
            fields["Resume Pages"] = rep["pages"]
    if (rep.get("ats") or {}).get("keyword_coverage_pct") is not None:
        fields["Keyword Coverage"] = rep["ats"]["keyword_coverage_pct"] / 100
    if files["cover_pdf"].exists():
        fields["Cover Letter"] = "Yes"
    meta["fields"] = fields
    config.write_json(meta_path, meta, indent=2)
    return str(meta_path)


def run(spec_path, overwrite=False, qa_only=False, qa_dir=None, paths=None):
    global FONT
    paths = paths or config.Paths()
    prof = config.profile(paths)
    FONT = (prof.get("style") or {}).get("font") or FONT
    spec_path = Path(spec_path)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    folder = package_folder(spec, paths)
    files = file_names(folder, prof)
    extra_links = re.findall(r"\]\((https?://[^)\s]+)\)", json.dumps(spec.get("resume") or {}))
    cl = spec.get("cover_letter")
    company = spec.get("company") or (folder.name.split(" - ")[0] if " - " in folder.name else None)
    title = f"{spec.get('company', '')} - {spec.get('title', '')}".strip(" -") or folder.name
    report = {"folder": str(folder)}

    if not qa_only:
        targets = []
        if spec.get("resume"):
            targets += [files["resume_docx"], files["resume_pdf"]]
        if spec.get("jd"):
            targets.append(files["jd_docx"])
        if cl:
            targets += [files["cover_docx"], files["cover_pdf"]]
        existing = [p for p in targets if p.exists()]
        if existing and not overwrite:
            raise SystemExit("Refusing to overwrite existing files (ask the user, or pass --overwrite "
                             "for files you created in this session):\n" + "\n".join(map(str, existing)))
        folder.mkdir(parents=True, exist_ok=True)
        try:
            if spec.get("resume"):
                build_resume(spec["resume"], prof, files["resume_docx"])
                report["pdf_method"] = export_pdf(files["resume_docx"], files["resume_pdf"])
            if spec.get("jd"):
                build_jd(spec["jd"], title, files["jd_docx"], spec_path.parent)
            if cl:
                build_cover(cl, prof, files["cover_docx"])
                report["pdf_method"] = export_pdf(files["cover_docx"], files["cover_pdf"])
        except ExportError as e:
            raise SystemExit(str(e)) from e

    qa_dir = Path(qa_dir) if qa_dir else Path(tempfile.gettempdir()) / "jhq_qa" / folder.name
    rep = {}
    if files["resume_pdf"].exists() and (spec.get("resume") or qa_only):
        rep = qa_resume(files["resume_docx"], files["resume_pdf"], qa_dir, prof, extra_links)
        if spec.get("ats"):
            from . import ats_check
            ats = ats_check.check(files["resume_pdf"], spec["ats"], prof=prof)
            rep["ats"] = {k: ats.get(k) for k in ("keyword_coverage_pct", "missing_required",
                                                  "keywords", "sections_found", "issues")}
            rep["warnings"] += [f"ATS parse: {i}" for i in ats["issues"]]
            if ats.get("missing_required"):
                rep["warnings"].append(f"ATS: required keywords missing {ats['missing_required']}: "
                                       "add only the ones that are true")
    report.update(rep)
    if cl and files["cover_pdf"].exists():
        report["cover_letter"] = qa_cover(cl, files["cover_pdf"], qa_dir, prof, company)
    if spec.get("company") and spec.get("title"):
        report["tracker_meta"] = write_meta(spec, folder, files, rep)
    report["files"] = {p.name: str(p) for k, p in files.items() if k != "meta" and p.exists()}
    return report


def main(argv=None):
    config.utf8_console()
    ap = argparse.ArgumentParser(prog="jhq build", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--overwrite", action="store_true")
    ap.add_argument("--qa-only", action="store_true")
    ap.add_argument("--qa-dir")
    a = ap.parse_args(argv)
    rep = run(a.spec, a.overwrite, a.qa_only, a.qa_dir)
    print(json.dumps(rep, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
